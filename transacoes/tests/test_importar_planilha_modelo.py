import io

from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from transacoes.xlsx_parse import PLANILHA_CABECALHOS


class ImportarPlanilhaModeloViewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth import get_user_model

        User = get_user_model()
        cls.user = User.objects.create_user(
            email="modelo_plan@test.local",
            password="secret",
            nome="Modelo",
        )

    def setUp(self):
        self.client = Client(enforce_csrf_checks=False)

    def test_requer_login(self):
        r = self.client.get(reverse("transacoes:importar_planilha_modelo"))
        self.assertEqual(r.status_code, 302)

    def test_get_retorna_xlsx_com_cabecalhos(self):
        self.client.force_login(self.user)
        r = self.client.get(reverse("transacoes:importar_planilha_modelo"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            r["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            row1 = next(ws.iter_rows(min_row=1, max_row=1, max_col=3, values_only=True))
            self.assertEqual(tuple(row1), PLANILHA_CABECALHOS)
        finally:
            wb.close()
