import io
from datetime import date, datetime
from decimal import Decimal

from django.test import SimpleTestCase
from openpyxl import Workbook

from transacoes.ofx_parse import LinhaOfx
from transacoes.xlsx_parse import PLANILHA_CABECALHOS, parse_planilha_padrao_bytes


def _bytes_workbook(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ParsePlanilhaPadraoBytesTests(SimpleTestCase):
    def test_cabecalho_errado(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["data", "valor", "descricao"])
        raw = _bytes_workbook(wb)
        with self.assertRaisesMessage(ValueError, "Cabeçalho inválido"):
            parse_planilha_padrao_bytes(raw)

    def test_arquivo_vazio(self):
        with self.assertRaisesMessage(ValueError, "Arquivo vazio."):
            parse_planilha_padrao_bytes(b"   \n")

    def test_parse_duas_linhas_e_ignora_vazia(self):
        wb = Workbook()
        ws = wb.active
        ws.append(list(PLANILHA_CABECALHOS))
        ws.append([date(2025, 3, 10), "Aluguel", -1200.00])
        ws.append([None, None, None])
        ws.append([datetime(2025, 3, 11, 8, 30), "Salário", 3000])
        raw = _bytes_workbook(wb)
        linhas = parse_planilha_padrao_bytes(raw)
        self.assertEqual(len(linhas), 2)
        self.assertEqual(
            linhas[0],
            LinhaOfx(
                data=date(2025, 3, 10),
                valor=Decimal("1200.00"),
                tipo="despesa",
                descricao="Aluguel",
                ofx_fitid=None,
            ),
        )
        self.assertEqual(linhas[1].data, date(2025, 3, 11))
        self.assertEqual(linhas[1].tipo, "receita")
        self.assertEqual(linhas[1].valor, Decimal("3000"))

    def test_data_string_iso(self):
        wb = Workbook()
        ws = wb.active
        ws.append(list(PLANILHA_CABECALHOS))
        ws.append(["2025-01-02", "X", 0])
        linhas = parse_planilha_padrao_bytes(_bytes_workbook(wb))
        self.assertEqual(linhas[0].tipo, "ajuste")
        self.assertEqual(linhas[0].valor, Decimal("0"))

    def test_descricao_vazia_vai_para_sem_descricao(self):
        wb = Workbook()
        ws = wb.active
        ws.append(list(PLANILHA_CABECALHOS))
        ws.append([date(2025, 1, 1), "", 10])
        linhas = parse_planilha_padrao_bytes(_bytes_workbook(wb))
        self.assertEqual(linhas[0].descricao, "Sem descrição")
