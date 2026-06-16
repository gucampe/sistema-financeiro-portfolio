from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from relatorios.export.base import ExportColumn, ExportDataset, ExportResult


def write_xlsx_dataset(dataset: ExportDataset) -> ExportResult:
  wb = Workbook(write_only=True)
  ws_data = wb.create_sheet("Dados")
  ws_meta = wb.create_sheet("Metadados")

  colunas = dataset.colunas
  header_font = Font(bold=True)
  ws_data.append([c.label for c in colunas])

  linhas_count = 0
  for row in dataset.linhas:
    ws_data.append([_format_cell(c, row.get(c.key)) for c in colunas])
    linhas_count += 1

  for key, value in dataset.metadados.items():
    if key == "nome_arquivo":
      continue
    ws_meta.append([str(key), str(value)])

  buf = io.BytesIO()
  wb.save(buf)
  content = buf.getvalue()
  nome = dataset.metadados.get("nome_arquivo", "export.xlsx")
  return ExportResult(
    content=content,
    nome_arquivo=nome if nome.endswith(".xlsx") else f"{nome}.xlsx",
    linhas_count=linhas_count,
    tamanho_bytes=len(content),
    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
  )


def _format_cell(col: ExportColumn, value: Any) -> Any:
  if col.formatter and value is not None:
    return col.formatter(value)
  if value is None:
    return ""
  return value
